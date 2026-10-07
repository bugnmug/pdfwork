"""Word documents for Word to PDF, built the way Word stores them: a report, a letter, a resume,
a document with Word's page features, one with real charts and a long table. All names, numbers
and companies are invented.

make-fixtures.py calls build(); run on its own: python3 word_fixtures.py [out_dir]
"""
import io, os, random, sys
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.opc.part import Part
from docx.opc.packuri import PackURI

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ------------------------------------------------------------------ pictures
def font(size, bold=False):
    try:
        return ImageFont.truetype("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf", size)
    except OSError:
        return ImageFont.load_default(size=size)


def png(img):
    b = io.BytesIO()
    img.save(b, "PNG")
    b.seek(0)
    return b


def logo_png():
    img = Image.new("RGB", (360, 120), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([4, 4, 116, 116], 18, fill=(31, 78, 121))
    d.polygon([(60, 30), (30, 90), (90, 90)], fill="white")
    d.text((134, 60), "Northwind", font=font(38, True), fill=(31, 78, 121), anchor="lm")
    return img


def chart_png():
    """A bar chart of revenue by region (the report's Figure 1)."""
    img = Image.new("RGB", (900, 450), "white")
    d = ImageDraw.Draw(img)
    x0, y0, x1, y1 = 96, 22, 878, 388
    top = 45.0
    f = font(18)
    for v in range(0, 41, 10):
        y = y1 - (y1 - y0) * v / top
        d.line([x0 - 8, y, x0, y], fill="black", width=2)
        d.text((x0 - 14, y), str(v), font=f, fill="black", anchor="rm")
    bars = [("North", 42.1, (31, 78, 121)), ("South", 35.6, (46, 117, 182)), ("East", 28.9, (157, 195, 230)), ("West", 31.4, (197, 90, 17))]
    band = (x1 - x0) / len(bars)
    for i, (name, v, col) in enumerate(bars):
        cx = x0 + band * (i + 0.5)
        d.rectangle([cx - band * 0.4, y1 - (y1 - y0) * v / top, cx + band * 0.4, y1], fill=col)
        d.line([cx, y1, cx, y1 + 8], fill="black", width=2)
        d.text((cx, y1 + 16), name, font=f, fill="black", anchor="mt")
    d.line([x0, y0, x0, y1], fill="black", width=2)
    d.line([x0, y1, x1, y1], fill="black", width=2)
    label = Image.new("RGB", (240, 30), "white")
    ImageDraw.Draw(label).text((120, 15), "Revenue (₹ lakh)", font=f, fill="black", anchor="mm")
    img.paste(label.rotate(90, expand=True), (22, (y0 + y1) // 2 - 120))
    return img


# ------------------------------------------------------------------ helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def field(par, instr, cached="1"):
    """A complex field (PAGE, NUMPAGES...) the way Word writes it."""
    r = par.add_run(); r._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
    r = par.add_run(); t = el("w:instrText", **{"xml:space": "preserve"}); t.text = f" {instr} "; r._r.append(t)
    r = par.add_run(); r._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    par.add_run(cached)
    r = par.add_run(); r._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))


def hyperlink(par, text, url, color="0563C1"):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = el("w:hyperlink", **{"r:id": rid})
    r = el("w:r"); rpr = el("w:rPr")
    rpr.append(el("w:rStyle", **{"w:val": "Hyperlink"}))
    rpr.append(el("w:color", **{"w:val": color})); rpr.append(el("w:u", **{"w:val": "single"}))
    r.append(rpr); t = el("w:t"); t.text = text; r.append(t); h.append(r); par._p.append(h)


def border(par, side="bottom", sz=8, color="1F4E79", space=1):
    ppr = par._p.get_or_add_pPr()
    b = ppr.find(qn("w:pBdr"))
    if b is None:
        b = el("w:pBdr"); ppr.append(b)
    b.append(el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": space, "w:color": color}))


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def cell_borders(table, color="BFBFBF", sz=4):
    b = el("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b.append(el(f"w:{side}", **{"w:val": "single", "w:sz": sz, "w:space": 0, "w:color": color}))
    table._tbl.tblPr.append(b)


def no_borders(table):
    b = el("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b.append(el(f"w:{side}", **{"w:val": "nil"}))
    table._tbl.tblPr.append(b)


def widths(table, cms):
    for row in table.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)


def run(par, text, bold=False, italic=False, size=None, color=None, font_name=None, underline=False, highlight=None):
    r = par.add_run(text)
    r.bold = bold or None; r.italic = italic or None; r.underline = underline or None
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if font_name: r.font.name = font_name
    if highlight: r._r.get_or_add_rPr().append(el("w:highlight", **{"w:val": highlight}))
    return r


def footnote(doc, par, text):
    """A real footnote: a footnotes.xml part, and the reference mark in the text."""
    main = doc.part
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes xmlns:w="{W_NS}">'
           '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
           '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
           f'<w:footnote w:id="1"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:footnote>'
           '</w:footnotes>')
    part = Part(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", xml.encode(), main.package)
    main.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes")
    r = par.add_run()
    r._r.get_or_add_rPr().append(el("w:vertAlign", **{"w:val": "superscript"}))
    r._r.append(el("w:footnoteReference", **{"w:id": 1}))


# ------------------------------------------------------------------ report
def report(out):
    """A4 report: header with a logo and a right tab, footer with Page X of Y, headings, a
    footnote, a link, lists, a shaded table, a picture, dot leaders and a landscape appendix."""
    doc = Document()
    st = doc.styles
    st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(11)
    st["Normal"].paragraph_format.space_after = Pt(8); st["Normal"].paragraph_format.line_spacing = 1.08
    for name, size, color in [("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2E75B6")]:
        s = st[name]; s.font.name = "Calibri Light"; s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color); s.font.bold = False
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.2)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    hp.add_run().add_picture(png(logo_png()), height=Cm(0.9))
    run(hp, "\tQuarterly Business Review  |  Q3 FY2026", size=9, color="595959")
    border(hp, "bottom", 6, "1F4E79", 4)
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    run(fp, "Northwind Traders Pvt. Ltd.  ·  Confidential", size=8, color="7F7F7F")
    run(fp, "\tPage ", size=8, color="7F7F7F"); field(fp, "PAGE"); run(fp, " of ", size=8, color="7F7F7F"); field(fp, "NUMPAGES", "3")
    for r in fp.runs: r.font.size = Pt(8)

    doc.add_paragraph(style="Title").add_run("Quarterly Business Review")
    doc.add_paragraph(style="Subtitle").add_run("July to September 2026, prepared for the board")
    doc.add_heading("1. Summary", level=1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p, "Revenue for the quarter reached ")
    run(p, "₹138.0 lakh", bold=True)
    run(p, ", up 18% on the same quarter last year and ahead of plan in every region except the East. Gross margin improved to ")
    run(p, "46.5%", bold=True, color="00B050")
    run(p, " as freight costs eased. Operating expenses grew more slowly than revenue for the third quarter in a row")
    footnote(doc, p, "Operating expenses exclude the one-time office move in August.")
    run(p, ". The full figures are published on the ")
    hyperlink(p, "investor portal", "https://example.com/investors/q3-2026")
    run(p, ".")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p, "Three matters need the board's decision this quarter: ")
    run(p, "the warehouse lease renewal", italic=True)
    run(p, ", the pricing change for wholesale customers, and the hiring plan for the second half. Each is described below with a recommendation. Items marked ")
    run(p, "urgent", highlight="yellow", bold=True)
    run(p, " must be settled before 15 October.")
    doc.add_heading("1.1 Highlights", level=2)
    for line in ["Online orders grew 31% and now make up a quarter of all sales.",
                 "Two new distributors signed in Pune and Coimbatore.",
                 "Customer complaints fell for the fourth quarter running."]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph("Returns processing moved in-house, cutting refund time from 9 days to 4.", style="List Bullet 2")
    doc.add_heading("1.2 Decisions needed", level=2)
    for line in ["Renew the Bhiwandi warehouse lease for five years at ₹4.2 lakh a month.",
                 "Raise wholesale prices by 4% from 1 November.",
                 "Approve 12 new roles, mostly in sales and support."]:
        doc.add_paragraph(line, style="List Number")

    doc.add_heading("2. Results by region", level=1)
    doc.add_paragraph("The table below compares this quarter with the same quarter last year. Figures are in ₹ lakh.")
    rows = [("Region", "Q3 FY2025", "Q3 FY2026", "Change"),
            ("North", "35.2", "42.1", "+19.6%"), ("South", "29.8", "35.6", "+19.5%"),
            ("East", "30.4", "28.9", "−4.9%"), ("West", "21.5", "31.4", "+46.0%"),
            ("Total", "116.9", "138.0", "+18.0%")]
    tb = doc.add_table(rows=len(rows), cols=4); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_borders(tb, "BFBFBF", 4)
    widths(tb, [5.0, 3.4, 3.4, 3.0])
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = ""
            par = c.paragraphs[0]; par.paragraph_format.space_after = Pt(0)
            if j: par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run(par, v, bold=(i == 0 or i == len(rows) - 1), color=("FFFFFF" if i == 0 else ("C00000" if v.startswith("−") else None)))
            if i == 0: shade(c, "1F4E79")
            elif i == len(rows) - 1: shade(c, "D9E2F3")
            elif i % 2 == 0: shade(c, "F2F2F2")
    doc.add_paragraph(style="Caption").add_run("Table 1. Revenue by region, ₹ lakh")
    doc.add_paragraph()
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(png(chart_png()), width=Cm(13))
    cap = doc.add_paragraph(style="Caption"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.add_run("Figure 1. Revenue by region this quarter")

    doc.add_heading("3. Outlook", level=1)
    for txt in ["The fourth quarter is usually the strongest, driven by festival demand. We expect revenue between ₹150 and ₹160 lakh if the wholesale price change goes ahead, and between ₹142 and ₹150 lakh if it does not.",
                "Freight costs are the main risk. A return to last year's rates would take about two points off the gross margin. We have fixed rates with our two largest carriers until March, which covers roughly 70% of volume."]:
        doc.add_paragraph(txt).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_heading("Key dates", level=2)
    for a, b in [("Board meeting", "15 October"), ("Price change takes effect", "1 November"), ("Q4 results", "12 January")]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run(p, a); run(p, "\t" + b, bold=True)

    new = doc.add_section(WD_SECTION.NEW_PAGE)
    new.orientation = WD_ORIENT.LANDSCAPE; new.page_width, new.page_height = Cm(29.7), Cm(21)
    doc.add_heading("Appendix A. Monthly detail", level=1)
    months = ["Jul", "Aug", "Sep"]
    hdr = ["Region"] + [f"{m} {k}" for m in months for k in ("Orders", "Revenue")] + ["Quarter"]
    tb = doc.add_table(rows=1, cols=len(hdr)); cell_borders(tb, "8EAADB", 4)
    for j, h in enumerate(hdr):
        c = tb.rows[0].cells[j]; c.text = ""; run(c.paragraphs[0], h, bold=True, size=9); shade(c, "D9E2F3")
    rnd = random.Random(4)
    for reg in ["North", "South", "East", "West"]:
        cells = tb.add_row().cells
        vals = [reg] + [str(rnd.randint(300, 900)) if k == 0 else f"{rnd.uniform(8, 16):.1f}" for m in months for k in range(2)] + [f"{rnd.uniform(28, 43):.1f}"]
        for j, v in enumerate(vals):
            cells[j].text = ""; par = cells[j].paragraphs[0]; run(par, v, size=9)
            if j: par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(os.path.join(out, "word-report.docx"))


# ------------------------------------------------------------------ letter
def letter(out):
    """US Letter, Times New Roman with a Georgia letterhead, first-line indents and a tab stop."""
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "MEHTA & IYER ASSOCIATES", bold=True, size=18, color="7B2C2C", font_name="Georgia")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "Chartered Accountants  ·  402 Lotus Court, Bandra West, Mumbai 400050  ·  +91 22 0000 1188", size=9, color="595959")
    border(p, "bottom", 12, "7B2C2C", 6)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(p, "6 October 2026")
    doc.add_paragraph()
    for line in ["Ms. Kavya Raman", "Finance Director", "Sundaram Looms Ltd.", "14 Mill Road, Coimbatore 641018"]:
        doc.add_paragraph(line)
    doc.add_paragraph()
    p = doc.add_paragraph(); run(p, "Subject: ", bold=True); run(p, "Engagement for the statutory audit, FY 2026–27", bold=True, underline=True)
    doc.add_paragraph()
    doc.add_paragraph("Dear Ms. Raman,")
    doc.add_paragraph()
    for b in ["Thank you for your letter of 28 September inviting us to act as statutory auditors of Sundaram Looms Ltd. for the financial year ending 31 March 2027. We are pleased to accept, subject to the terms set out below and to the approval of your shareholders at the annual general meeting.",
              "Our audit will be carried out in accordance with the Standards on Auditing issued by the Institute of Chartered Accountants of India. We will plan and perform the audit to obtain reasonable assurance about whether the financial statements are free from material misstatement, whether caused by fraud or error.",
              "The fees and timetable we propose are summarised below."]:
        p = doc.add_paragraph(b); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.5); p.paragraph_format.space_after = Pt(10)
    for a, b in [("Audit fee:", "₹6,50,000 plus GST"), ("Interim visit:", "January 2027"), ("Final visit:", "April to May 2027"), ("Report by:", "30 June 2027")]:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(2.2))
        run(p, a, bold=True); run(p, "\t" + b)
    doc.add_paragraph()
    p = doc.add_paragraph("Please sign and return the enclosed copy of this letter to confirm your agreement. We look forward to working with you."); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    doc.add_paragraph(); doc.add_paragraph("Yours sincerely,"); doc.add_paragraph(); doc.add_paragraph(); doc.add_paragraph()
    run(doc.add_paragraph(), "Rohan Mehta", bold=True)
    doc.add_paragraph("Partner, Membership No. 000000")
    doc.add_paragraph()
    p = doc.add_paragraph(); run(p, "Encl.: ", italic=True); run(p, "Copy of engagement terms", italic=True)
    doc.save(os.path.join(out, "word-letter.docx"))


# ------------------------------------------------------------------ resume
def resume(out):
    """Arial resume in a two-column table: a shaded sidebar, bullets, Hindi text, dates on right tabs."""
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.1
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.4)
    run(doc.add_paragraph(), "PRIYA NAIR", bold=True, size=24, color="0E5A6B")
    p = doc.add_paragraph(); run(p, "Product Manager  ·  Bengaluru  ·  priya.nair@example.com  ·  +91 98765 43210", size=9, color="404040")
    border(p, "bottom", 12, "0E5A6B", 4); p.paragraph_format.space_after = Pt(10)
    tb = doc.add_table(rows=1, cols=2); no_borders(tb); widths(tb, [5.6, 12.4])
    left, right = tb.rows[0].cells
    shade(left, "EAF2F4")

    def head(cell, text):
        par = cell.add_paragraph(); run(par, text.upper(), bold=True, size=10.5, color="0E5A6B")
        border(par, "bottom", 6, "0E5A6B", 1); par.paragraph_format.space_before = Pt(8); par.paragraph_format.space_after = Pt(4)

    left.paragraphs[0].text = ""
    head(left, "Skills")
    for s in ["Product discovery", "Roadmapping", "SQL and Looker", "A/B testing", "Stakeholder management"]:
        run(left.add_paragraph(style="List Bullet"), s, size=9.5)
    head(left, "Languages")
    for s in ["English (fluent)", "हिन्दी (fluent)", "Malayalam (native)", "Kannada (conversational)"]:
        run(left.add_paragraph(), s, size=9.5)
    head(left, "Education")
    run(left.add_paragraph(), "MBA, IIM Kozhikode", bold=True, size=9.5)
    run(left.add_paragraph(), "2016 – 2018", size=9, color="595959")
    run(left.add_paragraph(), "B.Tech, NIT Calicut", bold=True, size=9.5)
    run(left.add_paragraph(), "2010 – 2014", size=9, color="595959")

    right.paragraphs[0].text = ""
    head(right, "Profile")
    par = right.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(par, "Product manager with seven years in consumer fintech. Led the launch of a UPI payments app from zero to 4 million users, and rebuilt onboarding to halve drop-off. Comfortable with data, happiest with customers.")
    head(right, "Experience")
    jobs = [("Senior Product Manager", "PayLeaf", "2021 – present", ["Own the payments and rewards roadmap for 4M monthly users.", "Cut failed UPI transactions by 38% with smarter retry and bank routing.", "Run a team of 3 PMs, 2 designers and 18 engineers."]),
            ("Product Manager", "CrediNest", "2018 – 2021", ["Launched instant credit lines for small merchants; ₹120 crore disbursed in year one.", "Rebuilt the KYC flow; completion rose from 41% to 77%."]),
            ("Software Engineer", "Tessellate Software", "2014 – 2016", ["Built reporting tools for a European retail bank."])]
    for title, org, when, points in jobs:
        par = right.add_paragraph(); par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.tab_stops.add_tab_stop(Cm(12.0), WD_TAB_ALIGNMENT.RIGHT)
        run(par, title, bold=True); run(par, ", " + org); run(par, "\t" + when, size=9, color="595959")
        for pt in points:
            run(right.add_paragraph(style="List Bullet"), pt, size=9.5)
    doc.save(os.path.join(out, "word-resume.docx"))


# ------------------------------------------------------------------ features
def features(out):
    """Different first page, a DRAFT watermark, page borders, a contents list linked to its
    headings, a floating picture with square wrap, multilevel numbering, merged cells and a
    header row that repeats."""
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7); sec.left_margin = sec.right_margin = Cm(2.5)
    sec.different_first_page_header_footer = True
    sec.header.paragraphs[0].text = "Annual Plan 2027 · internal"
    sec.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sec.first_page_header.paragraphs[0].text = ""
    wm = parse_xml(f'''<w:r xmlns:w="{W_NS}" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e"><v:path textpathok="t" o:connecttype="custom"/><v:textpath on="t" fitshape="t"/></v:shapetype><v:shape id="PowerPlusWaterMarkObject" o:spid="_x0000_s2049" type="#_x0000_t136" style="position:absolute;margin-left:0;margin-top:0;width:412.4pt;height:206.2pt;rotation:315;z-index:-251657216;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" o:allowincell="f" fillcolor="silver" stroked="f"><v:fill opacity=".5"/><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="DRAFT"/></v:shape></w:pict></w:r>''')
    sec.header.paragraphs[0]._p.append(wm)
    sides = "".join(f'<w:{s} w:val="single" w:sz="12" w:space="24" w:color="1F4E79"/>' for s in ("top", "left", "bottom", "right"))
    sec._sectPr.insert(3, parse_xml(f'<w:pgBorders xmlns:w="{W_NS}" w:offsetFrom="page">{sides}</w:pgBorders>'))

    doc.add_heading("Annual Plan 2027", 0)

    def toc_entry(text, anchor, page):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        h = el("w:hyperlink", **{"w:anchor": anchor})
        for t in (text, "\t", page):
            r = el("w:r")
            if t == "\t":
                r.append(el("w:tab"))
            else:
                tt = el("w:t"); tt.text = t; r.append(tt)
            h.append(r)
        p._p.append(h)

    toc_entry("1. Goals", "_Toc1", "1")
    toc_entry("2. Plan by quarter", "_Toc2", "2")

    def heading(text, anchor, level=1):
        h = doc.add_heading("", level)
        h._p.append(el("w:bookmarkStart", **{"w:id": anchor[-1], "w:name": anchor}))
        h.add_run(text)
        h._p.append(el("w:bookmarkEnd", **{"w:id": anchor[-1]}))

    heading("1. Goals", "_Toc1")
    # A picture floated to the right of the column, text wrapping round it.
    p = doc.add_paragraph()
    r = p.add_run(); r.add_picture(png(chart_png()), width=Cm(6))
    drawing = r._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    anchor = parse_xml('<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="2" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV></wp:anchor>')
    for child in list(inline):
        if child.tag == qn("wp:extent"): anchor.append(child)
    ws = OxmlElement("wp:wrapSquare"); ws.set("wrapText", "bothSides"); anchor.append(ws)
    for child in list(inline):
        if child.tag in (qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")): anchor.append(child)
    drawing.remove(inline); drawing.append(anchor)
    p.add_run("We will grow revenue by a fifth while keeping costs flat. " * 6)
    # A multilevel list numbered 1., 1.1., (a).
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(f'''<w:abstractNum xmlns:w="{W_NS}" w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr></w:lvl>
<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="792" w:hanging="432"/></w:pPr></w:lvl>
<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="(%3)"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1224" w:hanging="432"/></w:pPr></w:lvl></w:abstractNum>'''))
    numbering.append(parse_xml(f'<w:num xmlns:w="{W_NS}" w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))

    def item(text, lvl):
        p = doc.add_paragraph(text)
        numPr = el("w:numPr"); numPr.append(el("w:ilvl", **{"w:val": lvl})); numPr.append(el("w:numId", **{"w:val": 90}))
        p._p.get_or_add_pPr().insert(0, numPr)

    for text, lvl in [("Grow revenue", 0), ("Open two new regions", 1), ("North by March", 2), ("East by June", 2), ("Raise prices by 4%", 1), ("Hold costs flat", 0), ("Renegotiate freight", 1)]:
        item(text, lvl)
    heading("2. Plan by quarter", "_Toc2")
    t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
    hdr = t.rows[0]
    for c, txt in zip(hdr.cells, ["Quarter", "Region", "Target", "Owner"]): c.text = txt
    hdr._tr.get_or_add_trPr().append(el("w:tblHeader"))
    for qi, q in enumerate(["Q1", "Q2", "Q3", "Q4"]):
        rows = []
        for reg in ["North", "South", "East", "West"]:
            row = t.add_row(); rows.append(row)
            row.cells[1].text = reg; row.cells[2].text = f"₹{10 + qi * 2 + len(reg)} lakh"; row.cells[3].text = "Asha Rao" if reg in ("North", "East") else "Vikram Shah"
        rows[0].cells[0].merge(rows[-1].cells[0]).text = q
    for i in range(30):
        doc.add_paragraph(f"Note {i + 1}: details of the plan for this line item, kept short so the table and notes run onto the next page.")
    doc.save(os.path.join(out, "word-features.docx"))


# ------------------------------------------------------------------ charts
C_NS = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def series(i, name, cats, vals, fmt="General", labels=False):
    cpts = "".join(f'<c:pt idx="{k}"><c:v>{c}</c:v></c:pt>' for k, c in enumerate(cats))
    vpts = "".join(f'<c:pt idx="{k}"><c:v>{v}</c:v></c:pt>' for k, v in enumerate(vals))
    dl = '<c:dLbls><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/></c:dLbls>' if labels else ""
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/><c:tx><c:strRef><c:f>Sheet1!$B$1</c:f><c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>{name}</c:v></c:pt></c:strCache></c:strRef></c:tx>{dl}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2</c:f><c:strCache><c:ptCount val="{len(cats)}"/>{cpts}</c:strCache></c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!$B$2</c:f><c:numCache><c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{vpts}</c:numCache></c:numRef></c:val></c:ser>')


AX_IDS = '<c:axId val="1"/><c:axId val="2"/>'
AXES = ('<c:catAx><c:axId val="1"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/><c:axPos val="b"/><c:crossAx val="2"/></c:catAx>'
        '<c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/><c:axPos val="l"/><c:majorGridlines/><c:numFmt formatCode="General" sourceLinked="1"/><c:crossAx val="1"/></c:valAx>')


def chart_space(title, body, legend="b"):
    t = f'<c:title><c:tx><c:rich><a:bodyPr/><a:p><a:r><a:t>{title}</a:t></a:r></a:p></c:rich></c:tx><c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/>'
    lg = f'<c:legend><c:legendPos val="{legend}"/><c:overlay val="0"/></c:legend>' if legend else ""
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><c:chartSpace {C_NS}><c:chart>{t}<c:plotArea><c:layout/>{body}</c:plotArea>{lg}<c:plotVisOnly val="1"/></c:chart></c:chartSpace>'


def charts(out):
    """Column, pie, line and stacked bar charts as Word stores them (chart parts, no pictures)."""
    q = ["Q1", "Q2", "Q3", "Q4"]
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    teams = ["Sales", "Support", "Product", "Ops"]
    parts = [
        chart_space("Revenue by quarter (₹ lakh)", f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>{series(0, "FY2025", q, [28.4, 30.1, 29.7, 33.2])}{series(1, "FY2026", q, [31.0, 35.4, 38.6, 41.9], labels=True)}<c:gapWidth val="150"/>{AX_IDS}</c:barChart>{AXES}'),
        chart_space("Sales mix", f'<c:pieChart><c:varyColors val="1"/>{series(0, "Share", ["Online", "Retail", "Wholesale", "Export"], [0.42, 0.31, 0.18, 0.09], fmt="0%", labels=True)}<c:firstSliceAng val="0"/></c:pieChart>', legend="r"),
        chart_space("Active customers", f'<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>{series(0, "North", months, [120, 132, 128, 150, 171, 180])}{series(1, "South", months, [90, 95, 110, 108, 120, 135])}<c:marker val="1"/>{AX_IDS}</c:lineChart>{AXES}'),
        chart_space("Headcount by team", f'<c:barChart><c:barDir val="bar"/><c:grouping val="stacked"/><c:varyColors val="0"/>{series(0, "Full time", teams, [24, 18, 12, 9])}{series(1, "Contract", teams, [6, 9, 3, 4])}<c:gapWidth val="80"/><c:overlap val="100"/>{AX_IDS}</c:barChart>{AXES}'),
    ]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.add_heading("Charts", 1)
    for i, xml in enumerate(parts, 1):
        part = Part(PackURI(f"/word/charts/chart{i}.xml"), "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", xml.encode(), doc.part.package)
        rid = doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart")
        w, h = int(Cm(14)), int(Cm(7.5))
        inline = parse_xml(f'<wp:inline distT="0" distB="0" distL="0" distR="0" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><wp:extent cx="{w}" cy="{h}"/><wp:docPr id="{100 + i}" name="Chart {i}"/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="{rid}"/></a:graphicData></a:graphic></wp:inline>')
        drawing = parse_xml(f'<w:drawing xmlns:w="{W_NS}"/>')
        drawing.append(inline)
        doc.add_paragraph().add_run()._r.append(drawing)
        doc.add_paragraph(f"Figure {i}.")
    doc.save(os.path.join(out, "word-charts.docx"))


# ------------------------------------------------------------------ long table
def longtable(out):
    """An 80-row table in a built-in table style whose header row repeats on every page."""
    doc = Document()
    doc.add_paragraph("Inventory list")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0]
    for c, txt in zip(hdr.cells, ["Item", "Quantity", "Location"]): c.text = txt
    hdr._tr.get_or_add_trPr().append(el("w:tblHeader"))
    for i in range(80):
        cells = t.add_row().cells
        cells[0].text = f"Item {i + 1}"
        cells[1].text = str(i * 37 % 200)
        cells[2].text = f"Aisle {i % 12}, shelf {i % 5}" + (" with a longer note that wraps onto a second line in this cell" if i % 7 == 0 else "")
    doc.save(os.path.join(out, "word-longtable.docx"))


def build(out):
    os.makedirs(out, exist_ok=True)
    for make in (report, letter, resume, features, charts, longtable):
        make(out)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures"))
    print("word fixtures written")
