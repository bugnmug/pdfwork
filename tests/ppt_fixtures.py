"""Write fixtures/deck-features.pptx: one slide per thing PowerPoint to PDF must get right.

1. Title slide (placeholders from the layout and master).
2. Shapes: solid, gradient and pattern fills, an outline, a rotated arrow, a flipped
   triangle, a group, an arrowed line and a shadow.
3. Bullets two levels deep, a numbered list and a hyperlink.
4. Pictures: a PNG with transparency and a WMF drawing (blue box, red disc, white "WMF").
5. A table in the default table style, with a merged cell.
6. A column chart with no formatting of its own.
7. A hidden slide (left out of the PDF).
8. A title that carries its own shape style while the master's title says "no fill",
   and a text box shrunk to half size by autofit.

Also fixtures/word-wmf.docx: a Word document with the same WMF picture.
"""
import copy, io, os, struct
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
q = lambda ns, tag: "{%s}%s" % (ns, tag)


def wmf():
    """A placeable WMF, 1 x 0.667 inch: a blue box on the left, a red disc on the right, "WMF" in white."""
    recs = []

    def rec(fn, *params, raw=b""):
        body = struct.pack("<" + "h" * len(params), *params) + raw
        if len(body) % 2:
            body += b"\0"
        recs.append(struct.pack("<IH", 3 + len(body) // 2, fn) + body)

    rec(0x020B, 0, 0)  # SetWindowOrg (y, x)
    rec(0x020C, 960, 1440)  # SetWindowExt (y, x)
    rec(0x02FC, raw=struct.pack("<HBBBBH", 0, 0x1F, 0x4E, 0x79, 0, 0))  # brush 0: blue
    rec(0x02FA, raw=struct.pack("<HhhBBBB", 5, 0, 0, 0, 0, 0, 0))  # pen 1: none
    rec(0x012D, 0)
    rec(0x012D, 1)
    rec(0x041B, 900, 700, 60, 60)  # Rectangle (bottom, right, top, left)
    rec(0x02FC, raw=struct.pack("<HBBBBH", 0, 0xC0, 0x00, 0x00, 0, 0))  # brush 2: red
    rec(0x012D, 2)
    rec(0x0418, 900, 1380, 60, 740)  # Ellipse
    rec(0x02FB, raw=struct.pack("<hhhhhBBBBBBBB", -180, 0, 0, 0, 700, 0, 0, 0, 0, 0, 0, 0, 0) + b"Arial".ljust(32, b"\0"))  # font 3
    rec(0x012D, 3)
    rec(0x0209, raw=struct.pack("<BBBB", 255, 255, 255, 0))  # text colour white
    rec(0x0102, 1)  # transparent background
    rec(0x0A32, 520, 180, 3, 0, raw=b"WMF\0")  # ExtTextOut (y, x, length, options, text)
    rec(0x0000)
    body = b"".join(recs)
    header = struct.pack("<HHHIHIH", 1, 9, 0x300, (18 + len(body)) // 2, 4, max(struct.unpack("<I", r[:4])[0] for r in recs), 0)
    placeable = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, 1440, 960, 1440, 0)
    check = 0
    for w in struct.unpack("<10H", placeable):
        check ^= w
    return placeable + struct.pack("<H", check) + header + body


def png_with_alpha():
    from PIL import Image, ImageDraw

    im = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
    ImageDraw.Draw(im).ellipse([20, 20, 180, 180], fill=(46, 125, 50, 255))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


prs = Presentation()
blank = prs.slide_layouts[6]

# 1. Title slide.
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "Feature Deck"
s.placeholders[1].text = "Shapes, text, tables and charts"

# 2. Shapes.
s = prs.slides.add_slide(blank)
box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(2), Inches(1))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0xC0, 0x00, 0x00)
box.line.fill.background()
box.text_frame.text = "Red box"
box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
oval = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(0.5), Inches(1.5), Inches(1.5))
oval.fill.gradient()
oval.fill.gradient_stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0x00)
oval.fill.gradient_stops[1].color.rgb = RGBColor(0x00, 0x70, 0xC0)
oval.line.color.rgb = RGBColor(0x20, 0x20, 0x20)
oval.line.width = Pt(2)
arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.2), Inches(0.6), Inches(2.2), Inches(1))
arrow.fill.solid()
arrow.fill.fore_color.rgb = RGBColor(0x00, 0xB0, 0x50)
arrow.rotation = 30
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(0.5), Inches(2.5), Inches(1.6), Inches(1.4))
tri.fill.solid()
tri.fill.fore_color.rgb = RGBColor(0xFF, 0xC0, 0x00)
tri._element.spPr.find(q(A, "xfrm")).set("flipV", "1")
star = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(2.6), Inches(2.4), Inches(1.6), Inches(1.6))
star.fill.patterned()
star.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
star.fill.fore_color.rgb = RGBColor(0x70, 0x30, 0xA0)
star.fill.back_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
grp = s.shapes.add_group_shape()
for i, c in enumerate((RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x8F, 0xAA, 0xDC))):
    g = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8 + i * 0.9), Inches(2.6), Inches(0.8), Inches(1.2))
    g.fill.solid()
    g.fill.fore_color.rgb = c
line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(4.6), Inches(4.5), Inches(4.6))
line.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
line.line.width = Pt(3)
ln = line._element.spPr.find(q(A, "ln"))
etree.SubElement(ln, q(A, "tailEnd"), type="triangle", w="med", len="med")
shade = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(4.3), Inches(2.2), Inches(1.2))
shade.fill.solid()
shade.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
shade.text_frame.text = "Shadowed"
shade.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 0, 0)
eff = etree.SubElement(shade._element.spPr, q(A, "effectLst"))
sh = etree.SubElement(eff, q(A, "outerShdw"), blurRad="50800", dist="76200", dir="2700000", algn="tl", rotWithShape="0")
clr = etree.SubElement(sh, q(A, "srgbClr"), val="000000")
etree.SubElement(clr, q(A, "alpha"), val="50000")

# 3. Bullets, numbering and a link.
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Bullets and links"
tf = s.placeholders[1].text_frame
tf.text = "First point"
for text, lvl in (("Detail one", 1), ("Detail two", 1), ("Second point with ", 0)):
    p = tf.add_paragraph()
    p.text = text
    p.level = lvl
run = tf.paragraphs[-1].add_run()
run.text = "a link"
run.hyperlink.address = "https://example.com/docs"
for text in ("Step alpha", "Step beta"):
    p = tf.add_paragraph()
    p.text = text
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "457200")
    pPr.set("indent", "-457200")
    etree.SubElement(pPr, q(A, "buAutoNum"), type="arabicPeriod")

# 4. Pictures.
s = prs.slides.add_slide(blank)
s.shapes.add_picture(io.BytesIO(png_with_alpha()), Inches(0.5), Inches(1), Inches(2.5), Inches(2.5))
s.shapes.add_picture(io.BytesIO(wmf()), Inches(4), Inches(1), Inches(4.5), Inches(3))
tb = s.shapes.add_textbox(Inches(0.5), Inches(5), Inches(8), Inches(0.6))
tb.text_frame.text = "A transparent PNG and a WMF drawing"

# 5. Table.
s = prs.slides.add_slide(blank)
rows = [("Region", "Q1", "Q2"), ("North", "120", "135"), ("South", "98", "110"), ("Total across regions", "", "")]
t = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(1), Inches(7), Inches(2.4)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        if v:
            t.cell(r, c).text = v
t.cell(3, 0).merge(t.cell(3, 2))

# 6. Chart.
s = prs.slides.add_slide(blank)
cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3", "Q4"]
cd.add_series("Revenue", (12, 18, 9, 24))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(0.8), Inches(8), Inches(5.5), cd).chart
ch.has_title = True
ch.chart_title.text_frame.text = "Quarterly revenue"

# 7. Hidden slide.
s = prs.slides.add_slide(blank)
s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "This slide is hidden"
s._element.set("show", "0")

# 8. A title with its own shape style over a master title with no fill; autofit at half size.
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Styled title"
title = s.shapes.title._element
style = etree.SubElement(title, q(P, "style"))
for tag, idx, val in (("lnRef", "1", "accent2"), ("fillRef", "1", "accent2"), ("effectRef", "0", "accent2"), ("fontRef", "minor", "lt1")):
    ref = etree.SubElement(style, q(A, tag), idx=idx)
    etree.SubElement(ref, q(A, "schemeClr"), val=val)
# p:style goes before p:txBody.
title.remove(style)
title.insert(list(title).index(title.find(q(P, "txBody"))), style)
master = prs.slide_masters[0]
for ph in master.placeholders:
    if ph.placeholder_format.type is not None and "TITLE" in str(ph.placeholder_format.type):
        spPr = ph._element.spPr
        for old in spPr.findall(q(A, "noFill")):
            spPr.remove(old)
        nofill = etree.Element(q(A, "noFill"))
        geom = spPr.find(q(A, "prstGeom"))
        (geom.addnext(nofill) if geom is not None else spPr.append(nofill))
small = s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1.5))
small.text_frame.text = "Shrunk to half size"
small.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
bodyPr = small.text_frame._txBody.find(q(A, "bodyPr"))
for old in list(bodyPr):
    bodyPr.remove(old)
etree.SubElement(bodyPr, q(A, "normAutofit"), fontScale="50000")

prs.save(os.path.join(OUT, "deck-features.pptx"))
print("wrote fixtures/deck-features.pptx")

# A Word document with the same WMF picture (python-docx can't add one: its PNG stand-in is swapped).
import zipfile
from docx import Document
from docx.shared import Inches as DInches
from PIL import Image

doc = Document()
doc.add_heading("Drawing in a metafile", 1)
doc.add_paragraph("The picture below is a WMF drawing: a blue box, a red disc and the letters WMF.")
stand_in = io.BytesIO()
Image.new("RGB", (72, 48), "white").save(stand_in, "PNG")
doc.add_picture(io.BytesIO(stand_in.getvalue()), width=DInches(4.5), height=DInches(3))
doc.add_paragraph("Text after the picture.")
raw = io.BytesIO()
doc.save(raw)
src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
out = zipfile.ZipFile(os.path.join(OUT, "word-wmf.docx"), "w", zipfile.ZIP_DEFLATED)
for item in src.infolist():
    data = src.read(item.filename)
    name = item.filename
    if name.startswith("word/media/") and name.endswith(".png"):
        name, data = name[:-4] + ".wmf", wmf()
    elif name == "word/_rels/document.xml.rels":
        data = data.replace(b".png", b".wmf")
    elif name == "[Content_Types].xml":
        data = data.replace(b"</Types>", b'<Default Extension="wmf" ContentType="image/x-wmf"/></Types>')
    out.writestr(name, data)
out.close()
print("wrote fixtures/word-wmf.docx")
