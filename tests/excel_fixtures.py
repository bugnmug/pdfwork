"""Excel workbooks for Excel to PDF, built the way spreadsheets are made: an invoice with a
logo and merged cells, a sales report fitted to one page wide with a chart sheet, an expense log
with gridlines and print titles, a sheet with an Excel table, conditional formats, a merged and a
rotated cell, a hidden row and a print area, and a program's export with no column widths set.
All names and numbers are invented. Formula results are filled in by a LibreOffice round trip,
as Excel saves them.

make-fixtures.py calls build(); run on its own: python3 excel_fixtures.py [out_dir]
"""
import datetime, io, os, random, shutil, subprocess, sys, tempfile, zipfile
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font as _Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageDraw, ImageFont


def Font(**kw):
    """Fonts named as Excel names them (openpyxl leaves the name out otherwise)."""
    kw.setdefault("name", "Calibri")
    kw.setdefault("size", 11)
    return _Font(**kw)


def logo():
    img = Image.new("RGB", (360, 120), "white")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([4, 4, 116, 116], 18, fill=(14, 90, 107))
    d.polygon([(60, 30), (30, 90), (90, 90)], fill="white")
    try:
        f = ImageFont.truetype("DejaVuSans-Bold.ttf", 38)
    except OSError:
        f = ImageFont.load_default(size=38)
    d.text((134, 60), "Tealbrook", font=f, fill=(14, 90, 107), anchor="lm")
    b = io.BytesIO()
    img.save(b, "PNG")
    b.seek(0)
    return b


thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color="0E5A6B")
INR = '"\u20b9"#,##0.00'


# ------------------------------------------------------------------ invoice
def invoice(out):
    wb = Workbook(); ws = wb.active; ws.title = "Invoice"
    for col, w in zip("ABCDEF", [6, 38, 10, 14, 10, 16]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:D1"); ws["A1"] = "Tealbrook Supplies Pvt. Ltd."
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="0E5A6B")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:D2"); ws["A2"] = "12 Harbour Road, Kochi 682001  ·  GSTIN 32ABCDE1234F1Z5"
    ws["A2"].font = Font(size=9, color="595959")
    img = XLImage(logo()); img.width, img.height = 150, 50; ws.add_image(img, "E1")
    ws["A4"] = "TAX INVOICE"; ws["A4"].font = Font(size=14, bold=True)
    ws["E4"] = "Invoice no."; ws["F4"] = "TB-2026-0418"
    ws["E5"] = "Date"; ws["F5"] = datetime.date(2026, 10, 6); ws["F5"].number_format = "dd-mmm-yyyy"
    ws["E6"] = "Due"; ws["F6"] = "=F5+30"; ws["F6"].number_format = "dd-mmm-yyyy"
    for c in ("E4", "E5", "E6"): ws[c].font = Font(bold=True, color="595959"); ws[c].alignment = Alignment(horizontal="right")
    for c in ("F4", "F5", "F6"): ws[c].alignment = Alignment(horizontal="right")
    ws["A6"] = "Bill to"; ws["A6"].font = Font(bold=True, color="595959")
    for i, line in enumerate(["Meridian Hotels Ltd.", "Attn: Purchase Department", "88 Marine Drive, Mumbai 400020"]):
        ws.cell(row=7 + i, column=1, value=line)
    hdr = ["#", "Description", "Qty", "Rate", "GST %", "Amount"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=11, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0E5A6B")
        c.alignment = Alignment(horizontal="center" if j != 2 else "left", vertical="center")
    ws.row_dimensions[11].height = 20
    items = [("Bath towels, 600 GSM, white", 240, 385), ("Bed sheets, king size, 300 thread count", 120, 1450),
             ("Pillow covers, set of two", 180, 320), ("Hand towels, 450 GSM", 300, 140),
             ("Laundry bags, printed with the hotel's name, drawstring", 400, 95)]
    for i, (desc, q, rate) in enumerate(items):
        r = 12 + i
        ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center", vertical="top")
        d = ws.cell(row=r, column=2, value=desc); d.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=q).number_format = "#,##0"
        ws.cell(row=r, column=4, value=rate).number_format = INR
        ws.cell(row=r, column=5, value=0.12 if i != 3 else 0.05).number_format = "0%"
        ws.cell(row=r, column=6, value=f"=C{r}*D{r}").number_format = INR
        for j in range(1, 7):
            ws.cell(row=r, column=j).border = Border(bottom=thin)
            if j > 2: ws.cell(row=r, column=j).alignment = Alignment(vertical="top", horizontal="right")
    last = 12 + len(items) - 1
    rows = [("Subtotal", f"=SUM(F12:F{last})"), ("GST", f"=SUMPRODUCT(F12:F{last},E12:E{last})"), ("Total", f"=F{last + 2}+F{last + 3}")]
    for k, (label, f) in enumerate(rows):
        r = last + 2 + k
        ws.cell(row=r, column=5, value=label).alignment = Alignment(horizontal="right")
        c = ws.cell(row=r, column=6, value=f); c.number_format = INR
        if label == "Total":
            for j in (5, 6):
                ws.cell(row=r, column=j).font = Font(bold=True, size=12)
                ws.cell(row=r, column=j).border = Border(top=med, bottom=Side(style="double", color="0E5A6B"))
    r = last + 6
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=6)
    n = ws.cell(row=r, column=1, value="Payment within 30 days to HDFC Bank, account 00000000001234, IFSC HDFC0000001. Goods once sold are taken back only if damaged in transit and reported within 7 days of delivery. Thank you for your business.")
    n.alignment = Alignment(wrap_text=True, vertical="top"); n.font = Font(size=9, italic=True, color="595959")
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    wb.save(os.path.join(out, "excel-invoice.xlsx"))

# ------------------------------------------------------------------ sales
def sales(out):
    wb = Workbook(); ws = wb.active; ws.title = "Sales 2026"
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    regions = ["North", "South", "East", "West", "Central", "North-East", "Islands", "Online"]
    ws["A1"] = "Monthly sales by region, 2026 (₹ thousand)"; ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:O1")
    ws["A2"] = "Region"; ws["B2"] = "Code"
    for j, m in enumerate(months): ws.cell(row=2, column=3 + j, value=m)
    ws.cell(row=2, column=15, value="Total")
    for j in range(1, 16):
        c = ws.cell(row=2, column=j); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="D9E2F3")
        c.border = Border(bottom=Side(style="thin", color="8EAADB")); c.alignment = Alignment(horizontal="center")
    rnd = random.Random(11)
    for i, reg in enumerate(regions):
        r = 3 + i
        ws.cell(row=r, column=1, value=reg)
        ws.cell(row=r, column=2, value=f"R{i + 1:02d}")
        for j in range(12):
            ws.cell(row=r, column=3 + j, value=rnd.randint(800, 4200)).number_format = "#,##0"
        ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})").number_format = "#,##0"
        ws.cell(row=r, column=15).font = Font(bold=True)
        if i % 2:
            for j in range(1, 16): ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="F2F2F2")
    tr = 3 + len(regions)
    ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
    for j in range(3, 16):
        col = get_column_letter(j)
        c = ws.cell(row=tr, column=j, value=f"=SUM({col}3:{col}{tr - 1})"); c.number_format = "#,##0"; c.font = Font(bold=True)
        c.border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    gr = tr + 1
    ws.cell(row=gr, column=1, value="Growth on 2025").font = Font(italic=True)
    for j in range(3, 15):
        c = ws.cell(row=gr, column=j, value=rnd.uniform(-0.08, 0.22)); c.number_format = "0.0%"
        c.font = Font(italic=True, color="C00000" if c.value < 0 else "00804A")
    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].hidden = True
    for j in range(3, 16): ws.column_dimensions[get_column_letter(j)].width = 9
    ws.column_dimensions["O"].width = 11
    ws.freeze_panes = "C3"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.center.text = "&A"; ws.oddFooter.left.text = "Tealbrook Supplies"; ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = "1:2"
    # A second sheet: a chart of the totals.
    cs = wb.create_sheet("Chart")
    cs["A1"] = "Region"; cs["B1"] = "Total"
    for i, reg in enumerate(regions):
        cs.cell(row=2 + i, column=1, value=reg); cs.cell(row=2 + i, column=2, value=f"='Sales 2026'!O{3 + i}").number_format = "#,##0"
    ch = BarChart(); ch.title = "Sales by region, 2026"; ch.y_axis.title = "₹ thousand"; ch.height = 9; ch.width = 18
    ch.add_data(Reference(cs, min_col=2, min_row=1, max_row=1 + len(regions)), titles_from_data=True)
    ch.set_categories(Reference(cs, min_col=1, min_row=2, max_row=1 + len(regions)))
    cs.add_chart(ch, "D2")
    wb.save(os.path.join(out, "excel-sales.xlsx"))

# ------------------------------------------------------------------ expenses
def expenses(out):
    wb = Workbook(); ws = wb.active; ws.title = "Expenses"
    hdr = ["Date", "Category", "Paid to", "Amount", "Notes"]
    for j, (h, w) in enumerate(zip(hdr, [12, 14, 22, 12, 40]), 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="404040")
        ws.column_dimensions[get_column_letter(j)].width = w
    rnd = random.Random(5)
    cats = ["Travel", "Meals", "Software", "Office", "Training"]
    vendors = ["IndiGo", "Uber", "Swiggy", "Atlassian", "Staples", "Coursera", "Taj Hotels", "Ola"]
    notes = ["", "", "Client visit to Pune, with the sales team", "", "Annual renewal for the design team; covers twelve seats and priority support", "", "Team lunch", ""]
    day = datetime.date(2026, 4, 1)
    for i in range(140):
        r = 2 + i
        day += datetime.timedelta(days=rnd.randint(0, 2))
        ws.cell(row=r, column=1, value=day).number_format = "dd-mmm-yyyy"
        ws.cell(row=r, column=2, value=rnd.choice(cats))
        ws.cell(row=r, column=3, value=rnd.choice(vendors))
        ws.cell(row=r, column=4, value=round(rnd.uniform(150, 48000), 2)).number_format = "#,##0.00"
        n = ws.cell(row=r, column=5, value=rnd.choice(notes)); n.alignment = Alignment(wrap_text=True, vertical="top")
        for j in range(1, 5): ws.cell(row=r, column=j).alignment = Alignment(vertical="top")
    ws.cell(row=143, column=3, value="Total").font = Font(bold=True)
    t = ws.cell(row=143, column=4, value="=SUM(D2:D141)"); t.number_format = "#,##0.00"; t.font = Font(bold=True)
    ws.print_options.gridLines = True
    ws.print_title_rows = "1:1"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.oddFooter.center.text = "&P / &N"
    wb.save(os.path.join(out, "excel-expenses.xlsx"))

# ------------------------------------------------------------------ features
def features(out):
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
    from openpyxl.chart import LineChart
    wb = Workbook(); ws = wb.active; ws.title = "Pipeline"
    ws["A1"] = "Sales pipeline, Q4 2026"; ws["A1"].font = Font(size=16, bold=True, color="1F3864")
    ws.merge_cells("A1:G1"); ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Figures in ₹ lakh unless stated. Deals marked at risk need a decision by Friday."
    ws["A2"].font = Font(italic=True, color="7F7F7F")
    hdr = ["Account", "Owner", "Stage", "Value", "Probability", "Weighted", "Close date"]
    rnd = random.Random(9)
    accounts = ["Meridian Hotels", "Saffron Retail", "Lotus Logistics", "Northstar Consulting", "Banyan Health", "Cobalt Mobility", "Juniper Foods", "Kestrel Steel", "Indigo Learning", "Opal Finance", "Tidewater Ports", "Vermilion Media"]
    stages = ["Discovery", "Proposal", "Negotiation", "Verbal yes", "At risk"]
    for j, h in enumerate(hdr, 1): ws.cell(row=4, column=j, value=h)
    for i, a in enumerate(accounts):
        r = 5 + i
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=rnd.choice(["Asha Rao", "Vikram Shah", "Neha Iyer"]))
        ws.cell(row=r, column=3, value=rnd.choice(stages))
        v = ws.cell(row=r, column=4, value=round(rnd.uniform(4, 180), 1)); v.number_format = '#,##,##0.0'
        pcell = ws.cell(row=r, column=5, value=round(rnd.uniform(0.1, 0.9), 2)); pcell.number_format = "0%"
        w = ws.cell(row=r, column=6, value=f"=D{r}*E{r}"); w.number_format = '_(* #,##0.0_);_(* \\(#,##0.0\\);_(* "-"??_);_(@_)'
        d = ws.cell(row=r, column=7, value=datetime.date(2026, 10, 1) + datetime.timedelta(days=rnd.randint(5, 85))); d.number_format = "d mmm yyyy"
    last = 4 + len(accounts)
    ws.cell(row=last + 1, column=1, value="Total")
    ws.cell(row=last + 1, column=4, value=f"=SUBTOTAL(109,D5:D{last})").number_format = '#,##,##0.0'
    ws.cell(row=last + 1, column=6, value=f"=SUBTOTAL(109,F5:F{last})").number_format = '_(* #,##0.0_);_(* \\(#,##0.0\\);_(* "-"??_);_(@_)'
    tab = Table(displayName="Pipeline", ref=f"A4:G{last + 1}", totalsRowCount=1)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    ws.conditional_formatting.add(f"C5:C{last}", CellIsRule(operator="equal", formula=['"At risk"'], fill=PatternFill("solid", bgColor="FFC7CE"), font=_Font(color="9C0006")))
    ws.conditional_formatting.add(f"E5:E{last}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws.conditional_formatting.add(f"D5:D{last}", DataBarRule(start_type="min", end_type="max", color="638EC6"))
    for col, w in zip("ABCDEFG", [24, 14, 13, 11, 12, 13, 13]): ws.column_dimensions[col].width = w
    # Notes: long text spilling right, a right-aligned label spilling left, wrapped text, a rotated label.
    n = last + 3
    ws.cell(row=n, column=1, value="Note: values exclude GST and assume the current price list; renewals are tracked on the second sheet.")
    ws.cell(row=n + 1, column=7, value="Prepared by the sales operations team").alignment = Alignment(horizontal="right")
    c = ws.cell(row=n + 2, column=1, value="Deals at risk are flagged in red; probability runs from red (low) to green (high), and the bars show deal size.")
    c.alignment = Alignment(wrap_text=True, vertical="top"); ws.merge_cells(start_row=n + 2, start_column=1, end_row=n + 2, end_column=4)
    ws.row_dimensions[n + 2].height = 45
    rl = ws.cell(row=n + 2, column=6, value="Reviewed"); rl.alignment = Alignment(text_rotation=90, horizontal="center", vertical="center"); rl.font = Font(bold=True, color="C00000")
    ws.row_dimensions[n + 4].hidden = True
    ws.cell(row=n + 4, column=1, value="HIDDEN ROW should not print")
    ws.oddHeader.right.text = "&D"; ws.oddFooter.center.text = "Page &P of &N"; ws.oddFooter.left.text = "&F"
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:G{n + 4}"
    # Second sheet: a long list with print titles and a chart.
    s2 = wb.create_sheet("Renewals")
    s2.append(["Month", "Renewals", "Churned", "Net"])
    for c2 in s2[1]: c2.font = Font(bold=True, color="FFFFFF"); c2.fill = PatternFill("solid", fgColor="2F5597")
    for i in range(60):
        m = datetime.date(2022, 1, 1) + datetime.timedelta(days=31 * i)
        ren = rnd.randint(20, 60); ch = rnd.randint(2, 15)
        s2.append([datetime.date(m.year, m.month, 1), ren, ch, ren - ch])
        s2.cell(row=2 + i, column=1).number_format = "mmm yyyy"
        s2.cell(row=2 + i, column=4).number_format = "0;[Red]-0"
    s2.print_title_rows = "1:1"
    lc = LineChart(); lc.title = "Net renewals by month"; lc.height = 7; lc.width = 15
    lc.add_data(Reference(s2, min_col=4, min_row=1, max_row=61), titles_from_data=True)
    s2.add_chart(lc, "F2")
    s2.oddFooter.right.text = "&A, page &P"
    wb.save(os.path.join(out, "excel-features.xlsx"))

def export(out):
    """A ledger written by a program: no column widths, dates and amounts in default-width columns."""
    wb = Workbook(); ws = wb.active; ws.title = "Ledger"
    ws.append(["Date", "Voucher", "Account", "Narration", "Debit", "Credit", "Balance", "Branch"])
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
    rnd = random.Random(1)
    bal = 0
    for i in range(120):
        d = round(rnd.uniform(1000, 250000), 2) if rnd.random() < 0.5 else 0
        cr = round(rnd.uniform(1000, 250000), 2) if d == 0 else 0
        bal += d - cr
        ws.append([datetime.date(2025, 4, 1) + datetime.timedelta(days=i // 4), f"JV-{i + 1:05d}", rnd.choice(["Sales", "Purchases", "Rent", "Salaries", "Bank charges"]), "Payment as per invoice " + str(rnd.randint(1000, 9999)), d, cr, bal, rnd.choice(["Pune", "Kochi", "Delhi"])])
        for col in (5, 6, 7): ws.cell(row=i + 2, column=col).number_format = "#,##0.00"
        ws.cell(row=i + 2, column=1).number_format = "dd-mm-yyyy"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    wb.save(os.path.join(out, "excel-export.xlsx"))



def build(out):
    os.makedirs(out, exist_ok=True)
    raw = tempfile.mkdtemp()
    for make in (invoice, sales, expenses, features):
        make(raw)
    export(out)
    # Formula results as Excel saves them: LibreOffice recalculates and writes the values.
    lo = tempfile.mkdtemp()
    for n in ("invoice", "sales", "expenses", "features"):
        name = f"excel-{n}.xlsx"
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", lo, os.path.join(raw, name)], check=True, capture_output=True)
        # LibreOffice drops a table's style when it saves: put the original table parts back.
        src = zipfile.ZipFile(os.path.join(raw, name))
        tables = {i.filename: src.read(i.filename) for i in src.infolist() if i.filename.startswith("xl/tables/")}
        done = zipfile.ZipFile(os.path.join(lo, name))
        items = [(i.filename, done.read(i.filename)) for i in done.infolist()]
        done.close()
        with zipfile.ZipFile(os.path.join(out, name), "w", zipfile.ZIP_DEFLATED) as z:
            for fname, data in items:
                z.writestr(fname, tables.get(fname, data))
    shutil.rmtree(raw, ignore_errors=True)
    shutil.rmtree(lo, ignore_errors=True)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures"))
    print("excel fixtures written")
